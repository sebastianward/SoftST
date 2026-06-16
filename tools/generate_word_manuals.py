from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path.cwd()
ANNOTATED = ROOT / "manual_assets" / "annotated"


def build_docx(filename: str, title: str, sections: list[dict]) -> None:
    document = Document()

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(11)

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_paragraph.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(22)
    title_run.font.color.rgb = RGBColor(47, 29, 89)

    intro = document.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.add_run("Manual de uso con capturas y referencias visuales por funcionalidad.")

    for section in sections:
        document.add_heading(section["heading"], level=1)

        for paragraph in section.get("paragraphs", []):
            document.add_paragraph(paragraph)

        for bullet in section.get("bullets", []):
            document.add_paragraph(bullet, style="List Bullet")

        if section.get("note"):
            note = document.add_paragraph()
            note_run = note.add_run(f"Nota: {section['note']}")
            note_run.italic = True

        if section.get("image"):
            image_path = ANNOTATED / section["image"]
            document.add_picture(str(image_path), width=Inches(6.7))
            caption = document.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption_run = caption.add_run(section.get("caption", ""))
            caption_run.italic = True

    document.save(ROOT / filename)


docs = {
    "MANUAL_USUARIO_ADMIN.docx": {
        "title": "Manual de Usuario Admin",
        "sections": [
            {
                "heading": "Alcance",
                "paragraphs": [
                    "El perfil admin tiene acceso completo sobre el area activa. Puede cambiar entre Servicio Tecnico y Grafica, administrar trabajadores, revisar configuraciones, operar camionetas y revisar notificaciones."
                ],
            },
            {
                "heading": "1. Iniciar sesion",
                "image": "admin_login.png",
                "caption": "Campos a completar para ingresar al sistema.",
                "bullets": [
                    "1: escribe tu correo o usuario.",
                    "2: escribe tu contrasena.",
                    "3: haz clic en Entrar.",
                ],
            },
            {
                "heading": "2. Navegacion principal",
                "image": "admin_dashboard.png",
                "caption": "Accesos principales del perfil admin.",
                "bullets": [
                    "1: Nuevo ingreso para crear un registro nuevo.",
                    "2: Trabajadores para enrolar, editar y cambiar perfil entre user y operator.",
                    "3: Configuracion para ajustes generales de correo y plazos.",
                    "4: Camionetas para asignaciones, mantencion e historial del modulo.",
                    "5: Notificaciones para revisar alertas pendientes.",
                ],
            },
            {
                "heading": "3. Gestion de trabajadores",
                "image": "admin_workers.png",
                "caption": "Mantenedor de trabajadores con cambio de perfil.",
                "bullets": [
                    "1: buscador por nombre o correo.",
                    "2: selector de Perfil para cambiar entre user y operator.",
                    "3: Guardar cambios para aplicar nombre, correo, area, perfil o estado.",
                    "4: Restablecer clave para generar una nueva clave temporal.",
                ],
                "note": "El perfil admin no se asigna desde esta pantalla; solo permite alternar entre user y operator.",
            },
            {
                "heading": "4. Camionetas - Asignaciones",
                "image": "admin_camionetas_assignments.png",
                "caption": "Asignacion diaria de camionetas.",
                "bullets": [
                    "1: subpestana Asignaciones.",
                    "2: buscador con sugerencia para escoger al trabajador.",
                    "3: Guardar para asignar o reasignar la camioneta.",
                ],
            },
            {
                "heading": "5. Camionetas - Mantencion",
                "image": "admin_camionetas_manage.png",
                "caption": "Creacion, edicion y horario de reseteo.",
                "bullets": [
                    "1: subpestana Mantencion.",
                    "2: formulario para crear una camioneta nueva o un reemplazo.",
                    "3: hora diaria de reseteo automatico del turno.",
                    "4: edicion rapida del listado existente.",
                ],
            },
            {
                "heading": "6. Camionetas - Historial",
                "image": "admin_camionetas_history.png",
                "caption": "Historial detallado de movimientos del modulo.",
                "bullets": [
                    "1: subpestana Historial.",
                    "2: tabla con eventos del turno y cambios historicos.",
                ],
            },
            {
                "heading": "7. Notificaciones",
                "image": "admin_notifications.png",
                "caption": "Revision y cierre de alertas.",
                "bullets": [
                    "1: listado de notificaciones activas o revisadas.",
                    "2: Marcar revisada para cerrar visualmente una alerta pendiente.",
                ],
            },
        ],
    },
    "MANUAL_USUARIO_OPERADOR.docx": {
        "title": "Manual de Usuario Operador",
        "sections": [
            {
                "heading": "Alcance",
                "paragraphs": [
                    "El perfil operator opera ingresos y camionetas dentro de Servicio Tecnico. Puede editar campos operativos del historial y acceder al mantenedor e historial de camionetas."
                ],
            },
            {
                "heading": "1. Iniciar sesion",
                "image": "operator_login.png",
                "caption": "Ingreso al sistema.",
                "bullets": [
                    "1: escribe tu correo o usuario.",
                    "2: escribe tu contrasena.",
                    "3: haz clic en Entrar.",
                ],
            },
            {
                "heading": "2. Navegacion principal",
                "image": "operator_dashboard.png",
                "caption": "Accesos principales del perfil operator.",
                "bullets": [
                    "1: Ingresos para revisar y actualizar el historial.",
                    "2: Nuevo ingreso para registrar un equipo nuevo.",
                    "3: Camionetas para asignaciones, mantencion e historial del turno.",
                ],
            },
            {
                "heading": "3. Historial de ingresos",
                "image": "operator_entries.png",
                "caption": "Campos operativos editables del historial.",
                "bullets": [
                    "1: buscador general.",
                    "2: campos operativos editables de la fila.",
                    "3: acceso a modales para texto largo como comentario o detalle.",
                    "4: Reimprimir para volver a emitir la etiqueta del ingreso.",
                ],
                "note": "Como operador puedes editar estado, codigo SAP, comentario, tarea final, cotizacion y OC.",
            },
            {
                "heading": "4. Camionetas - Asignaciones",
                "image": "operator_camionetas_assignments.png",
                "caption": "Asignacion operativa de camionetas.",
                "bullets": [
                    "1: subpestana Asignaciones.",
                    "2: buscador con sugerencias para elegir al trabajador.",
                    "3: Guardar para asignar o reasignar.",
                ],
            },
            {
                "heading": "5. Camionetas - Mantencion",
                "image": "operator_camionetas_manage.png",
                "caption": "Mantenedor de camionetas y reemplazos.",
                "bullets": [
                    "1: subpestana Mantencion.",
                    "2: formulario para agregar camionetas o reemplazos.",
                    "3: horario de reseteo automatico del turno.",
                ],
            },
            {
                "heading": "6. Camionetas - Historial",
                "image": "operator_camionetas_history.png",
                "caption": "Revision de movimientos del turno.",
                "bullets": [
                    "1: subpestana Historial.",
                    "2: tabla de movimientos del modulo.",
                ],
            },
        ],
    },
    "MANUAL_USUARIO_USER.docx": {
        "title": "Manual de Usuario User",
        "sections": [
            {
                "heading": "Alcance",
                "paragraphs": [
                    "El perfil user registra ingresos, consulta el historial en modo lectura y puede asignar camionetas dentro de Servicio Tecnico."
                ],
            },
            {
                "heading": "1. Iniciar sesion",
                "image": "user_login.png",
                "caption": "Ingreso al sistema.",
                "bullets": [
                    "1: escribe tu correo o usuario.",
                    "2: escribe tu contrasena.",
                    "3: haz clic en Entrar.",
                ],
            },
            {
                "heading": "2. Navegacion principal",
                "image": "user_dashboard.png",
                "caption": "Accesos principales del perfil user.",
                "bullets": [
                    "1: Nuevo ingreso para registrar un equipo.",
                    "2: Ingresos para consultar historial.",
                    "3: Camionetas para asignar o reasignar camionetas.",
                ],
            },
            {
                "heading": "3. Registrar un nuevo ingreso",
                "image": "user_new_entry.png",
                "caption": "Formulario principal para registrar equipos.",
                "bullets": [
                    "1: Razon social, uno de los campos obligatorios.",
                    "2: Reporte de cliente, obligatorio.",
                    "3: Ingresado por, selecciona el trabajador usando el buscador con sugerencias.",
                    "4: carga imagenes desde galeria o camara.",
                    "5: Guardar ingreso para registrar el equipo.",
                ],
            },
            {
                "heading": "4. Consultar historial",
                "image": "user_entries.png",
                "caption": "Consulta en modo lectura.",
                "bullets": [
                    "1: buscador general para filtrar resultados.",
                    "2: sliders superior e inferior para revisar columnas laterales.",
                    "3: filas del historial solo para consulta.",
                ],
            },
            {
                "heading": "5. Camionetas - Asignaciones",
                "image": "user_camionetas_assignments.png",
                "caption": "Asignacion de camionetas para el turno.",
                "bullets": [
                    "1: buscador con sugerencias para elegir trabajador.",
                    "2: Guardar para asignar o reasignar la camioneta.",
                ],
            },
        ],
    },
}


for filename, payload in docs.items():
    build_docx(filename, payload["title"], payload["sections"])
